"""Step 5: export scripts from the vault as Markdown, CSV, JSON or Excel."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .agent_script import rebuild_result
from .report import render_script

COLUMNS = ["id", "title", "niche", "status", "quality", "topic", "angle", "framework", "framework_evidence", "tags", "notes", "audio_hook",
           "text_hook", "video_hook", "spoken_script", "caption", "hashtags", "cta", "needs_from_you", "source_reel", "created_at"]
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value) -> str:
    """A spreadsheet will run a cell that starts with = + - @ as a formula. Prefix those so pasted script text stays text."""
    s = "" if value is None else str(value)
    return "'" + s if s.startswith(FORMULA_STARTS) else s


def flatten(row: dict) -> dict:
    d, hooks = row["data"], row["data"].get("hooks", {})
    fw = d.get("framework", {})
    return {"id": row["id"], "title": row["title"], "niche": row["niche"], "status": row["status"], "quality": row["quality"],
            "topic": row["topic"], "angle": row["angle"], "framework": fw.get("name", row["framework_id"]), "framework_evidence": row["framework_level"],
            "tags": ", ".join(row["tags"]), "notes": row["notes"], "audio_hook": hooks.get("audio", ""), "text_hook": hooks.get("text", ""),
            "video_hook": hooks.get("video", ""), "spoken_script": " ".join(b.get("spoken", "") for b in d.get("beats", []) if b.get("spoken")),
            "caption": d.get("caption", ""), "hashtags": " ".join(d.get("hashtags", [])), "cta": (d.get("cta") or {}).get("text", ""),
            "needs_from_you": "; ".join(d.get("needs_from_you", [])), "source_reel": row["source_reel_id"], "created_at": row["created_at"]}


def beats_of(row: dict) -> list[dict]:
    return [{"script_id": row["id"], "title": row["title"], "time": (f"{b['start_s']:g}-{b['end_s']:g}s" if b.get("start_s") is not None and b.get("end_s") is not None else ""),
             "beat": b.get("beat", ""), "say": b.get("spoken", ""), "on_screen_text": b.get("on_screen_text", ""), "show": b.get("visual", ""), "sound": b.get("audio", "")}
            for b in row["data"].get("beats", [])]


def export_scripts(rows: list[dict], fmt: str, path: Path, cfg: dict, date: str = "") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "json":
        path.write_text(json.dumps(rows, indent=2, default=str), encoding="utf-8")
    elif fmt == "md":
        parts = [f"<!-- script {r['id']} · status {r['status']} · tags {', '.join(r['tags'])} -->\n" + render_script(rebuild_result(r), date, False, cfg) +
                 (f"\n\n**Notes:** {r['notes']}" if r["notes"] else "") for r in rows]
        path.write_text("\n\n---\n\n".join(parts), encoding="utf-8")
    elif fmt == "csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:          # utf-8-sig so Excel reads accents and emoji correctly
            w = csv.DictWriter(fh, fieldnames=COLUMNS, quoting=csv.QUOTE_ALL)
            w.writeheader()
            for r in rows:
                w.writerow({k: safe_cell(v) for k, v in flatten(r).items()})
    elif fmt == "xlsx":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
        except ImportError as exc:
            raise RuntimeError("Excel export needs openpyxl: pip install openpyxl (or export csv and open it in Sheets).") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Scripts"
        ws.append(COLUMNS)
        for r in rows:
            ws.append([safe_cell(v) for v in flatten(r).values()])
        shots = wb.create_sheet("Shot lists")
        heads = ["script_id", "title", "time", "beat", "say", "on_screen_text", "show", "sound"]
        shots.append(heads)
        for r in rows:
            for b in beats_of(r):
                shots.append([safe_cell(b[h]) for h in heads])
        for sheet in (ws, shots):
            for c in sheet[1]:
                c.font = Font(bold=True)
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = min(max(len(str(col[0].value)) + 4, 14), 48)
            for row in sheet.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.freeze_panes = "A2"
        wb.save(path)
    else:
        raise ValueError("format must be one of: md, csv, json, xlsx")
    return path
